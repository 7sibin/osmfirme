from __future__ import annotations

from dataclasses import replace

from openpyxl import load_workbook

from osm_businesses import COLUMNS, Row

from webapp.export import FOUND_COLUMNS, build_workbook, export_filename
from webapp.results import ResultFilters


def make_row(name="Pekara Sunce") -> Row:
    return Row(
        osm_type="node", osm_id=1, name=name, category="bakery", place="Nis",
        street="Obrenoviceva", housenumber="10", postcode="18000",
        phone="+38118111222", phone_alt="", website="https://example.rs",
        email="pekara@example.rs", facebook="", instagram="",
        opening_hours="Mo-Fr 08:00-20:00", lat=43.3152812, lon=21.8944804,
        category_key="shop",
    )


def load(rows, **kwargs):
    kwargs.setdefault("area_label", "Nis")
    kwargs.setdefault("filters", ResultFilters())
    return load_workbook(build_workbook(rows, **kwargs))


def test_workbook_has_a_data_sheet_and_an_info_sheet():
    book = load([make_row()])
    assert book.sheetnames == ["Firme", "Info"]


def test_header_row_is_serbian_and_matches_the_column_count():
    sheet = load([make_row()])["Firme"]
    headers = [cell.value for cell in sheet[1]]
    assert len(headers) == len(COLUMNS) + len(FOUND_COLUMNS)
    assert headers[2] == "Naziv"


def test_what_the_search_found_is_its_own_labelled_column():
    """Not folded into `Sajt`: it did not come from the map and ODbL does not cover it."""
    found = replace(make_row("Pekara Trpkovic"), website="", found_website="https://trpkovic.rs",
                    found_confidence="strong", found_source="search")
    sheet = load([found])["Firme"]
    headers = [cell.value for cell in sheet[1]]
    values = dict(zip(headers, [cell.value for cell in sheet[2]]))
    assert not values["Sajt"]  # openpyxl writes an empty cell as None
    assert values["Sajt (pretraga)"] == "https://trpkovic.rs"
    assert values["Pouzdanost"] == "sigurno"
    assert values["Odakle"] == "rezultat pretrage"


def test_a_weak_hit_is_labelled_as_needing_a_look():
    found = replace(make_row("A"), found_website="https://x.rs", found_confidence="weak")
    sheet = load([found])["Firme"]
    values = dict(zip([c.value for c in sheet[1]], [c.value for c in sheet[2]]))
    assert values["Pouzdanost"] == "za proveru"


def test_what_reading_the_site_found_is_also_its_own_column():
    read = replace(make_row("Pekara Trpkovic"), email="", phone="",
                   found_email="info@trpkovic.rs", found_phone="018512345",
                   contact_status="ok")
    sheet = load([read])["Firme"]
    values = dict(zip([c.value for c in sheet[1]], [c.value for c in sheet[2]]))
    assert not values["Email"]
    assert values["Email (sa sajta)"] == "info@trpkovic.rs"
    assert values["Telefon (sa sajta)"] == "018512345"
    assert values["Sajt procitan"] == "da"


def test_a_dead_site_says_so_in_words():
    dead = replace(make_row("A"), contact_status="dead")
    sheet = load([dead])["Firme"]
    values = dict(zip([c.value for c in sheet[1]], [c.value for c in sheet[2]]))
    assert values["Sajt procitan"] == "sajt ne radi"


def test_an_unchecked_row_says_so_rather_than_looking_like_a_miss():
    sheet = load([make_row("A")])["Firme"]
    values = dict(zip([c.value for c in sheet[1]], [c.value for c in sheet[2]]))
    assert values["Pouzdanost"] == "nije provereno"
    assert values["Sajt procitan"] == "nije citan"


def test_one_row_per_business():
    sheet = load([make_row("A"), make_row("B")])["Firme"]
    assert sheet.max_row == 3  # header + two rows


def test_phone_keeps_its_leading_plus_as_text():
    sheet = load([make_row()])["Firme"]
    phone_column = COLUMNS.index("phone") + 1
    cell = sheet.cell(row=2, column=phone_column)
    assert cell.value == "+38118111222"
    assert cell.data_type == "s"


def test_website_and_email_are_hyperlinks():
    sheet = load([make_row()])["Firme"]
    website = sheet.cell(row=2, column=COLUMNS.index("website") + 1)
    email = sheet.cell(row=2, column=COLUMNS.index("email") + 1)
    assert website.hyperlink is not None
    assert email.hyperlink.target == "mailto:pekara@example.rs"


def test_empty_website_gets_no_hyperlink():
    row = make_row()
    blank = Row(**{**row.as_output_dict(), "website": "", "email": ""}, category_key="shop")
    sheet = load([blank])["Firme"]
    assert sheet.cell(row=2, column=COLUMNS.index("website") + 1).hyperlink is None


def test_coordinates_stay_numbers():
    sheet = load([make_row()])["Firme"]
    assert sheet.cell(row=2, column=COLUMNS.index("lat") + 1).value == 43.3152812


def test_first_row_is_frozen_and_autofilter_is_on():
    sheet = load([make_row()])["Firme"]
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref is not None


def test_info_sheet_carries_the_area_and_the_odbl_attribution():
    text = "\n".join(
        str(cell.value) for row in load([make_row()])["Info"].iter_rows() for cell in row if cell.value
    )
    assert "Nis" in text
    assert "OpenStreetMap" in text


def test_info_sheet_reports_the_applied_filters():
    filters = ResultFilters(categories=["bakery"], require_contact=True, q="pek")
    text = "\n".join(
        str(cell.value)
        for row in load([make_row()], filters=filters)["Info"].iter_rows()
        for cell in row
        if cell.value
    )
    assert "bakery" in text
    assert "pek" in text


def test_empty_result_still_produces_a_valid_workbook():
    sheet = load([])["Firme"]
    assert sheet.max_row == 1


def test_filename_is_slugged_and_dated():
    name = export_filename("Ниш, Град Ниш")
    assert name.startswith("firme-")
    assert name.endswith(".xlsx")
    assert " " not in name


def test_info_sheet_reports_the_website_filter():
    text = "\n".join(
        str(cell.value)
        for row in load([make_row()], filters=ResultFilters(website="no"))["Info"].iter_rows()
        for cell in row
        if cell.value
    )
    assert "bez sajta" in text
