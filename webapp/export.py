"""Turn a filtered list of rows into an .xlsx workbook in memory."""

from __future__ import annotations

import re
import unicodedata
from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from osm_businesses import COLUMNS, Row

from webapp.results import FOUND_STRONG, FOUND_WEAK, WEBSITE_NO, WEBSITE_YES, ResultFilters

HEADERS_SR: dict[str, str] = {
    "osm_type": "OSM tip",
    "osm_id": "OSM ID",
    "name": "Naziv",
    "category": "Kategorija",
    "place": "Mesto",
    "street": "Ulica",
    "housenumber": "Broj",
    "postcode": "Postanski broj",
    "phone": "Telefon",
    "phone_alt": "Drugi telefon",
    "website": "Sajt",
    "email": "Email",
    "facebook": "Facebook",
    "instagram": "Instagram",
    "opening_hours": "Radno vreme",
    "lat": "Geo. sirina",
    "lon": "Geo. duzina",
}

#: Appended after the OSM columns. Kept separate and labelled rather than folded
#: into `website`, because these did not come from the map: ODbL covers the OSM
#: half of the sheet and nothing else, and a weak hit is the user's call to make.
FOUND_COLUMNS: tuple[str, ...] = (
    "found_website", "found_confidence", "found_source",
    "found_email", "found_phone", "contact_status",
)

HEADERS_FOUND: dict[str, str] = {
    "found_website": "Sajt (pretraga)",
    "found_confidence": "Pouzdanost",
    "found_source": "Odakle",
    "found_email": "Email (sa sajta)",
    "found_phone": "Telefon (sa sajta)",
    "contact_status": "Sajt procitan",
}

#: What reading the site came to, in words.
CONTACT_LABELS: dict[str, str] = {
    "ok": "da",
    "none": "da, nema kontakta",
    "dead": "sajt ne radi",
    "": "nije citan",
}

#: Turned into words: a spreadsheet full of `strong` helps nobody.
FOUND_LABELS: dict[str, str] = {
    FOUND_STRONG: "sigurno",
    FOUND_WEAK: "za proveru",
    "none": "nije nadjen",
    "": "nije provereno",
}

SOURCE_LABELS: dict[str, str] = {
    "search": "rezultat pretrage",
    "mention": "pomenut u tekstu",
    "social": "drustvena mreza",
    "": "",
}

ATTRIBUTION = (
    "Podaci (c) OpenStreetMap contributors, ODbL. "
    "https://www.openstreetmap.org/copyright"
)

MAX_COLUMN_WIDTH = 45
TEXT_COLUMNS = ("phone", "phone_alt", "postcode", "housenumber", "found_phone")
LINK_FONT = Font(color="0563C1", underline="single")


def _slug(text: str) -> str:
    ascii_text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")
    slug = re.sub(r"[^a-zA-Z0-9]+", "-", ascii_text).strip("-").lower()
    return slug or "oblast"


def export_filename(area_label: str) -> str:
    return f"firme-{_slug(area_label)}-{date.today().isoformat()}.xlsx"


def build_workbook(rows: list[Row], *, area_label: str, filters: ResultFilters) -> BytesIO:
    book = Workbook()
    sheet = book.active
    sheet.title = "Firme"

    sheet_columns = (*COLUMNS, *FOUND_COLUMNS)
    headers = {**HEADERS_SR, **HEADERS_FOUND}

    for index, column in enumerate(sheet_columns, start=1):
        cell = sheet.cell(row=1, column=index, value=headers[column])
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center")

    for row_index, row in enumerate(rows, start=2):
        values = {
            **row.as_output_dict(),
            "found_website": row.found_website,
            "found_confidence": FOUND_LABELS.get(row.found_confidence, row.found_confidence),
            "found_source": SOURCE_LABELS.get(row.found_source, row.found_source),
            "found_email": row.found_email,
            "found_phone": row.found_phone,
            "contact_status": CONTACT_LABELS.get(row.contact_status, row.contact_status),
        }
        for column_index, column in enumerate(sheet_columns, start=1):
            value = values[column]
            cell = sheet.cell(row=row_index, column=column_index)
            if column in ("lat", "lon"):
                cell.value = float(value)
                cell.number_format = "0.0000000"
            elif column in TEXT_COLUMNS:
                cell.value = str(value)
                cell.number_format = "@"  # text, so Excel keeps the leading +
            else:
                cell.value = value
            if value and column in ("website", "found_website"):
                cell.hyperlink = str(value)
                cell.font = LINK_FONT
            elif value and column in ("email", "found_email"):
                cell.hyperlink = f"mailto:{value}"
                cell.font = LINK_FONT

    sheet.freeze_panes = "A2"
    last_column = get_column_letter(len(sheet_columns))
    sheet.auto_filter.ref = f"A1:{last_column}{max(1, len(rows) + 1)}"
    _fit_columns(sheet, rows, sheet_columns, headers)

    _write_info_sheet(book, rows=rows, area_label=area_label, filters=filters)

    buffer = BytesIO()
    book.save(buffer)
    buffer.seek(0)
    return buffer


def _fit_columns(sheet, rows: list[Row], columns, headers) -> None:
    for index, column in enumerate(columns, start=1):
        widest = len(headers[column])
        for row in rows:
            widest = max(widest, len(str(getattr(row, column, ""))))
        sheet.column_dimensions[get_column_letter(index)].width = min(widest + 2, MAX_COLUMN_WIDTH)


def _write_info_sheet(book: Workbook, *, rows: list[Row], area_label: str, filters: ResultFilters) -> None:
    info = book.create_sheet("Info")
    applied = []
    if filters.categories:
        applied.append("kategorije: " + ", ".join(sorted(filters.categories)))
    if filters.require_contact:
        applied.append("samo sa kontaktom")
    if filters.website == WEBSITE_YES:
        applied.append("samo sa sajtom")
    elif filters.website == WEBSITE_NO:
        applied.append("samo bez sajta")
    if filters.q.strip():
        applied.append(f"pretraga: {filters.q.strip()}")
    if filters.collapse:
        applied.append("lanci sazeti na jedan red")
    if filters.commercial_only:
        applied.append("bez banaka, posta i javnih ustanova")
    if filters.hide_found:
        applied.append("sakriveni oni kojima je pretraga nasla sajt")

    found = sum(1 for row in rows if row.found_website)
    emails = sum(1 for row in rows if row.found_email)
    dead = sum(1 for row in rows if row.contact_status == "dead")

    lines: list[tuple[str, object]] = [
        ("Oblast", area_label),
        ("Datum", date.today().isoformat()),
        ("Broj firmi", len(rows)),
        ("Filteri", "; ".join(applied) or "bez filtera"),
        ("Izvor", ATTRIBUTION),
        ("Sajtovi iz pretrage", found),
        ("Email adresa sa sajtova", emails),
        ("Sajtova koji ne rade", dead),
        (
            "Napomena",
            "Kolone koje pocinju sa 'Sajt (pretraga)' nisu iz OpenStreetMap-a nego iz "
            "veb pretrage i sa sajtova samih firmi, i nisu pokrivene ODbL licencom. "
            "Redove oznacene 'za proveru' pogledaj pre nego sto ih koristis. "
            "'Sajt ne radi' znaci da firma ima sajt u OSM-u koji vise ne odgovara.",
        ),
    ]
    for row_index, (label, value) in enumerate(lines, start=1):
        info.cell(row=row_index, column=1, value=label).font = Font(bold=True)
        info.cell(row=row_index, column=2, value=value)
    info.column_dimensions["A"].width = 16
    info.column_dimensions["B"].width = 80
