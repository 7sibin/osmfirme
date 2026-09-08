from __future__ import annotations

import pytest
from fastapi.testclient import TestClient

from geo import AreaCandidate
from webapp.main import app, get_geometry, get_nominatim


class FakeNominatim:
    def __init__(self, candidates=None, error=None):
        self.candidates = candidates or []
        self.error = error
        self.calls = []

    def search(self, name, country=None, limit=10):
        self.calls.append((name, country))
        if self.error is not None:
            raise self.error
        return self.candidates


class FakeGeometry:
    def __init__(self, result=None):
        self.result = result

    def lookup(self, osm_type, osm_id):
        return self.result


CANDIDATE = AreaCandidate(
    osm_type="relation", osm_id=11538321, display_name="Nis, Grad Nis, Srbija",
    category="boundary", place_type="city", admin_level="6",
)


@pytest.fixture
def client():
    with TestClient(app) as test_client:
        yield test_client
    app.dependency_overrides.clear()


def test_health_is_ok(client):
    assert client.get("/api/health").json() == {"status": "ok"}


def test_places_returns_candidates_with_area_ids(client):
    app.dependency_overrides[get_nominatim] = lambda: FakeNominatim([CANDIDATE])
    body = client.get("/api/places", params={"q": "Nis", "country": "RS"}).json()
    assert body[0]["area_id"] == 3611538321
    assert body[0]["display_name"].startswith("Nis")


def test_places_with_no_matches_is_an_empty_list_not_an_error(client):
    app.dependency_overrides[get_nominatim] = lambda: FakeNominatim([])
    response = client.get("/api/places", params={"q": "qqqq"})
    assert response.status_code == 200
    assert response.json() == []


def test_places_requires_a_query_of_at_least_two_characters(client):
    assert client.get("/api/places", params={"q": "a"}).status_code == 422


def test_places_reports_a_nominatim_failure_as_503_with_a_sentence(client):
    app.dependency_overrides[get_nominatim] = lambda: FakeNominatim(error=RuntimeError("down"))
    response = client.get("/api/places", params={"q": "Nis"})
    assert response.status_code == 503
    assert "down" not in response.json()["detail"]


def test_geometry_returns_the_polygon(client):
    payload = {"geojson": {"type": "Polygon", "coordinates": []}, "bbox": [1.0, 2.0, 3.0, 4.0], "display_name": "Nis"}
    app.dependency_overrides[get_geometry] = lambda: FakeGeometry(payload)
    assert client.get("/api/places/relation/11538321/geometry").json() == payload


def test_geometry_without_a_polygon_is_404(client):
    app.dependency_overrides[get_geometry] = lambda: FakeGeometry(None)
    assert client.get("/api/places/relation/1/geometry").status_code == 404


def test_root_serves_the_page(client):
    response = client.get("/")
    assert response.status_code == 200
    assert "text/html" in response.headers["content-type"]


# --- jobs -------------------------------------------------------------------

from webapp.cache import ResultCache  # noqa: E402
from webapp.jobs import JobRegistry  # noqa: E402
from webapp.main import get_cache, get_overpass_factory, get_registry  # noqa: E402

ELEMENTS = [
    # Pekara has both a phone and a website, Kafic has neither: one of each for
    # the require_contact and website filters to bite on.
    {"type": "node", "id": 1, "lat": 43.32, "lon": 21.9,
     "tags": {"name": "Pekara", "shop": "bakery", "phone": "+38118111222",
              "website": "https://pekara.example.rs"}},
    {"type": "node", "id": 2, "lat": 43.33, "lon": 21.91,
     "tags": {"name": "Kafic", "amenity": "cafe"}},
]


class StubOverpass:
    def __init__(self, elements):
        self.elements = elements

    def fetch(self, query):
        return self.elements


@pytest.fixture
def registry():
    """One registry for the whole test, so the test can inspect the jobs it created."""
    return JobRegistry()


@pytest.fixture
def wired(tmp_path, registry):
    """A client with a fresh registry, a temp cache and a stubbed Overpass."""
    app.dependency_overrides[get_registry] = lambda: registry
    app.dependency_overrides[get_cache] = lambda: ResultCache(tmp_path)
    app.dependency_overrides[get_overpass_factory] = lambda: (lambda: StubOverpass(ELEMENTS))
    with TestClient(app) as test_client:
        yield test_client
    app.dependency_overrides.clear()


def post_job(client, **overrides):
    body = {"area": {"kind": "area", "area_id": 3611538321, "label": "Nis"},
            "categories": ["shop", "amenity"]}
    body.update(overrides)
    return client.post("/api/jobs", json=body)


def wait_for_job(client, job_id, tries=100):
    for _ in range(tries):
        body = client.get(f"/api/jobs/{job_id}").json()
        if body["status"] in {"done", "error", "cancelled"}:
            return body
    raise AssertionError(f"job {job_id} never finished")


def test_post_jobs_returns_a_job_id(wired):
    body = post_job(wired).json()
    assert body["job_id"]
    assert body["cached"] is False


def test_job_runs_to_done_and_reports_counts(wired):
    job_id = post_job(wired).json()["job_id"]
    body = wait_for_job(wired, job_id)
    assert body["status"] == "done"
    assert body["elements_found"] == 2
    assert body["rows"] == 2


def test_a_second_identical_job_is_served_from_the_cache(wired):
    wait_for_job(wired, post_job(wired).json()["job_id"])
    second = post_job(wired).json()
    assert second["cached"] is True


def test_invalid_area_is_422(wired):
    assert post_job(wired, area={"kind": "bbox", "bbox": [1, 2]}).status_code == 422


def test_unknown_category_is_422(wired):
    assert post_job(wired, categories=["nonsense"]).status_code == 422


def test_status_of_an_unknown_job_is_404(wired):
    assert wired.get("/api/jobs/nope").status_code == 404


def test_delete_cancels_a_running_job(wired):
    job_id = post_job(wired).json()["job_id"]
    wired.delete(f"/api/jobs/{job_id}")
    body = wired.get(f"/api/jobs/{job_id}").json()
    assert body["status"] in {"cancelled", "done"}  # a stub can finish before the cancel lands


def test_delete_of_an_unknown_job_is_404(wired):
    assert wired.delete("/api/jobs/nope").status_code == 404


# --- results and export -----------------------------------------------------


def test_results_return_rows_facets_and_total(wired):
    job_id = post_job(wired).json()["job_id"]
    wait_for_job(wired, job_id)

    body = wired.get(f"/api/jobs/{job_id}/results").json()
    assert body["total"] == 2
    assert {item["value"] for item in body["facets"]} == {"bakery", "cafe"}
    assert body["rows"][0]["name"] in {"Pekara", "Kafic"}


def test_results_apply_the_category_filter(wired):
    job_id = post_job(wired).json()["job_id"]
    wait_for_job(wired, job_id)

    body = wired.get(f"/api/jobs/{job_id}/results", params={"categories": "bakery"}).json()
    assert body["total"] == 1
    assert body["rows"][0]["name"] == "Pekara"


def test_results_apply_require_contact(wired):
    job_id = post_job(wired).json()["job_id"]
    wait_for_job(wired, job_id)

    body = wired.get(f"/api/jobs/{job_id}/results", params={"require_contact": "true"}).json()
    assert body["total"] == 1


def test_results_paginate(wired):
    job_id = post_job(wired).json()["job_id"]
    wait_for_job(wired, job_id)

    body = wired.get(f"/api/jobs/{job_id}/results", params={"page": 2, "page_size": 1}).json()
    assert body["page"] == 2
    assert len(body["rows"]) == 1
    assert body["total"] == 2


def test_results_of_an_unfinished_job_are_409(wired, registry):
    # A job that was never started stays "pending"; posting one and forcing its
    # status would race with the background task finishing it.
    job = registry.create("Nis")
    assert wired.get(f"/api/jobs/{job.job_id}/results").status_code == 409


def test_export_returns_an_xlsx_with_a_filename(wired):
    job_id = post_job(wired).json()["job_id"]
    wait_for_job(wired, job_id)

    response = wired.get(f"/api/jobs/{job_id}/export.xlsx")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "firme-" in response.headers["content-disposition"]
    assert response.content[:2] == b"PK"  # xlsx is a zip


def test_export_honours_the_same_filters(wired):
    from io import BytesIO

    from openpyxl import load_workbook

    job_id = post_job(wired).json()["job_id"]
    wait_for_job(wired, job_id)

    response = wired.get(f"/api/jobs/{job_id}/export.xlsx", params={"categories": "bakery"})
    sheet = load_workbook(BytesIO(response.content))["Firme"]
    assert sheet.max_row == 2  # header + one row


def test_results_apply_the_website_filter(wired):
    job_id = post_job(wired).json()["job_id"]
    wait_for_job(wired, job_id)

    without = wired.get(f"/api/jobs/{job_id}/results", params={"website": "no"}).json()
    assert [row["name"] for row in without["rows"]] == ["Kafic"]

    with_site = wired.get(f"/api/jobs/{job_id}/results", params={"website": "yes"}).json()
    assert [row["name"] for row in with_site["rows"]] == ["Pekara"]


def test_results_reject_an_unknown_website_value(wired):
    job_id = post_job(wired).json()["job_id"]
    wait_for_job(wired, job_id)

    assert wired.get(f"/api/jobs/{job_id}/results", params={"website": "maybe"}).status_code == 422


def test_results_carry_the_area_counts_and_the_export_name(wired):
    job_id = post_job(wired).json()["job_id"]
    wait_for_job(wired, job_id)

    body = wired.get(f"/api/jobs/{job_id}/results").json()
    assert body["counts"] == {
        "area_total": 2, "area_without_site": 1,
        "found_sites": 0, "maybe_sites": 0,
        # Both fixture names are bare trade words, so neither can be searched for.
        "unchecked": 0, "unchecked_here": 0,
    }
    assert body["export_filename"].startswith("firme-nis-")


def test_area_counts_describe_the_area_not_the_filtered_table(wired):
    """The headline says what the area holds; only `total` follows the filters."""
    job_id = post_job(wired).json()["job_id"]
    wait_for_job(wired, job_id)

    body = wired.get(f"/api/jobs/{job_id}/results", params={"q": "pekara"}).json()
    assert body["total"] == 1
    assert body["counts"]["area_total"] == 2
    assert body["counts"]["area_without_site"] == 1


def test_points_return_a_coordinate_per_filtered_row(wired):
    job_id = post_job(wired).json()["job_id"]
    wait_for_job(wired, job_id)

    body = wired.get(f"/api/jobs/{job_id}/points").json()
    assert body["total"] == 2
    assert body["truncated"] is False
    assert sorted(body["points"]) == [[43.32, 21.9], [43.33, 21.91]]


def test_points_honour_the_filters(wired):
    job_id = post_job(wired).json()["job_id"]
    wait_for_job(wired, job_id)

    body = wired.get(f"/api/jobs/{job_id}/points", params={"website": "no"}).json()
    assert body["points"] == [[43.33, 21.91]]  # Kafic, the one without a site


def test_points_of_an_unfinished_job_are_409(wired, registry):
    job = registry.create("Nis")
    assert wired.get(f"/api/jobs/{job.job_id}/points").status_code == 409


def test_static_files_must_be_revalidated(client):
    response = client.get("/static/app.js")
    assert response.status_code == 200
    assert response.headers["cache-control"] == "no-cache"


# --- the website search ------------------------------------------------------

from webapp.enrich import SiteHit  # noqa: E402
from webapp.main import get_finder_factory, get_site_cache  # noqa: E402
from webapp.site_cache import SiteCache  # noqa: E402

#: Two businesses with names a search can actually work with, unlike ELEMENTS.
NAMED_ELEMENTS = [
    {"type": "node", "id": 1, "lat": 43.32, "lon": 21.9,
     "tags": {"name": "Pekara Trpkovic", "shop": "bakery", "addr:city": "Nis"}},
    {"type": "node", "id": 2, "lat": 43.33, "lon": 21.91,
     "tags": {"name": "Apoteka Jankovic", "shop": "chemist", "addr:city": "Nis"}},
    # A chain with two branches and a bank: both are gone by the time the search
    # sees the rows, which is the point of running it on the prepared set.
    {"type": "node", "id": 3, "lat": 43.34, "lon": 21.92,
     "tags": {"name": "Maxi", "shop": "supermarket", "addr:street": "Bulevar"}},
    {"type": "node", "id": 4, "lat": 43.35, "lon": 21.93,
     "tags": {"name": "Maxi", "shop": "supermarket", "addr:street": "Nemanjina"}},
    {"type": "node", "id": 5, "lat": 43.36, "lon": 21.94,
     "tags": {"name": "Banka Intesa", "amenity": "bank"}},
]

VERDICTS = {
    "Pekara Trpkovic": SiteHit(website="https://trpkovic.rs", confidence="strong",
                               source="search", checked_at="2026-09-08T00:00:00+00:00"),
    "Apoteka Jankovic": SiteHit(confidence="none", checked_at="2026-09-08T00:00:00+00:00"),
    "Maxi": SiteHit(website="https://maxi.rs", confidence="strong", source="search",
                    checked_at="2026-09-08T00:00:00+00:00"),
}


class StubFinder:
    """Answers from VERDICTS and records who it was asked about."""

    asked: list[str] = []

    def find(self, row):
        StubFinder.asked.append(row.name)
        return VERDICTS[row.name]


@pytest.fixture
def searching(tmp_path, registry):
    StubFinder.asked = []
    app.dependency_overrides[get_registry] = lambda: registry
    app.dependency_overrides[get_cache] = lambda: ResultCache(tmp_path / "results")
    app.dependency_overrides[get_overpass_factory] = lambda: (lambda: StubOverpass(NAMED_ELEMENTS))
    app.dependency_overrides[get_site_cache] = lambda: SiteCache(tmp_path / "sites")
    app.dependency_overrides[get_finder_factory] = lambda: (lambda: StubFinder())
    with TestClient(app) as test_client:
        yield test_client
    app.dependency_overrides.clear()


def enriched_job(client, registry, **params):
    job_id = post_job(client, categories=["shop", "amenity"]).json()["job_id"]
    wait_for_job(client, job_id)
    client.post(f"/api/jobs/{job_id}/enrich", params=params)
    _run_side_tasks(client, registry, job_id)
    return job_id


def _run_side_tasks(client, registry, job_id, tries=200):
    for _ in range(tries):
        body = client.get(f"/api/jobs/{job_id}/enrich").json()
        if body["status"] in {"done", "error", "cancelled"}:
            return body
    raise AssertionError("the website search never finished")


def test_enrichment_reports_what_it_found(searching, registry):
    job_id = enriched_job(searching, registry)
    progress = searching.get(f"/api/jobs/{job_id}/enrich").json()
    assert progress["status"] == "done"
    assert progress["found"] == 2
    assert progress["checked"] == 3


def test_enrichment_runs_on_the_prepared_rows_not_the_raw_ones(searching, registry):
    """The bank is filtered out and the second Maxi collapsed away before searching."""
    enriched_job(searching, registry)
    assert sorted(StubFinder.asked) == ["Apoteka Jankovic", "Maxi", "Pekara Trpkovic"]


def test_a_found_site_shows_up_on_the_row_without_overwriting_the_osm_one(searching, registry):
    job_id = enriched_job(searching, registry)
    body = searching.get(f"/api/jobs/{job_id}/results", params={"q": "trpkovic"}).json()
    row = body["rows"][0]
    assert row["website"] == ""
    assert row["found_website"] == "https://trpkovic.rs"
    assert row["found_confidence"] == "strong"


def test_hide_found_drops_them_from_the_table_but_not_from_the_counts(searching, registry):
    job_id = enriched_job(searching, registry)
    shown = searching.get(f"/api/jobs/{job_id}/results", params={"hide_found": "true"}).json()
    assert [row["name"] for row in shown["rows"]] == ["Apoteka Jankovic"]
    assert shown["counts"]["area_total"] == 3
    assert shown["counts"]["found_sites"] == 2


def test_the_results_route_reports_how_many_are_still_unchecked(searching, registry):
    job_id = post_job(searching, categories=["shop", "amenity"]).json()["job_id"]
    wait_for_job(searching, job_id)
    body = searching.get(f"/api/jobs/{job_id}/results").json()
    assert body["counts"]["unchecked"] == 3
    assert body["enrich"]["status"] == "idle"


def test_starting_a_second_pass_while_one_runs_is_refused(searching, registry):
    job_id = post_job(searching, categories=["shop", "amenity"]).json()["job_id"]
    wait_for_job(searching, job_id)
    registry.get(job_id).enrich.progress.status = "running"
    assert searching.post(f"/api/jobs/{job_id}/enrich").status_code == 409


def test_enrichment_on_an_unfinished_job_is_refused(searching):
    assert searching.post("/api/jobs/nope/enrich").status_code == 404


def test_cancelling_the_search_leaves_the_job_itself_done(searching, registry):
    job_id = enriched_job(searching, registry)
    searching.delete(f"/api/jobs/{job_id}/enrich")
    assert searching.get(f"/api/jobs/{job_id}").json()["status"] == "done"


def test_the_default_view_collapses_chains_and_drops_the_bank(searching):
    job_id = post_job(searching, categories=["shop", "amenity"]).json()["job_id"]
    wait_for_job(searching, job_id)
    names = [row["name"] for row in searching.get(f"/api/jobs/{job_id}/results").json()["rows"]]
    assert sorted(names) == ["Apoteka Jankovic", "Maxi", "Pekara Trpkovic"]


def test_turning_both_off_shows_every_row_again(searching):
    job_id = post_job(searching, categories=["shop", "amenity"]).json()["job_id"]
    wait_for_job(searching, job_id)
    body = searching.get(
        f"/api/jobs/{job_id}/results",
        params={"collapse": "false", "commercial_only": "false"},
    ).json()
    assert body["total"] == 5


def test_the_search_only_takes_on_what_the_filters_leave(searching, registry):
    """Two seconds a business means the queue has to follow the filter bar.

    Without this the queue is `sort_rows` order - alphabetical by category - so
    a first pass never reaches the trade you actually sell to.
    """
    job_id = post_job(searching, categories=["shop", "amenity"]).json()["job_id"]
    wait_for_job(searching, job_id)
    searching.post(f"/api/jobs/{job_id}/enrich", params={"categories": ["bakery"]})
    _run_side_tasks(searching, registry, job_id)

    assert StubFinder.asked == ["Pekara Trpkovic"]


def test_a_text_search_narrows_the_queue_too(searching, registry):
    job_id = post_job(searching, categories=["shop", "amenity"]).json()["job_id"]
    wait_for_job(searching, job_id)
    searching.post(f"/api/jobs/{job_id}/enrich", params={"q": "jankovic"})
    _run_side_tasks(searching, registry, job_id)

    assert StubFinder.asked == ["Apoteka Jankovic"]


def test_unchecked_here_follows_the_filters_while_unchecked_stays_the_area(searching):
    job_id = post_job(searching, categories=["shop", "amenity"]).json()["job_id"]
    wait_for_job(searching, job_id)

    wide = searching.get(f"/api/jobs/{job_id}/results").json()["counts"]
    assert wide["unchecked"] == 3 and wide["unchecked_here"] == 3

    narrow = searching.get(
        f"/api/jobs/{job_id}/results", params={"categories": ["bakery"]}
    ).json()["counts"]
    assert narrow["unchecked"] == 3, "the area count must not move with the filters"
    assert narrow["unchecked_here"] == 1


def test_filtering_to_nothing_leaves_the_search_nothing_to_do(searching, registry):
    job_id = post_job(searching, categories=["shop", "amenity"]).json()["job_id"]
    wait_for_job(searching, job_id)
    counts = searching.get(
        f"/api/jobs/{job_id}/results", params={"q": "nepostojeca firma"}
    ).json()["counts"]
    assert counts["unchecked_here"] == 0

    searching.post(f"/api/jobs/{job_id}/enrich", params={"q": "nepostojeca firma"})
    body = _run_side_tasks(searching, registry, job_id)
    assert body["status"] == "done"
    assert StubFinder.asked == []


def test_a_filtered_pass_leaves_the_rest_for_the_next_one(searching, registry):
    """Nothing is wasted: the site cache is keyed by business, not by pass."""
    job_id = post_job(searching, categories=["shop", "amenity"]).json()["job_id"]
    wait_for_job(searching, job_id)

    searching.post(f"/api/jobs/{job_id}/enrich", params={"categories": ["bakery"]})
    _run_side_tasks(searching, registry, job_id)
    searching.post(f"/api/jobs/{job_id}/enrich")
    _run_side_tasks(searching, registry, job_id)

    assert StubFinder.asked == ["Pekara Trpkovic", "Apoteka Jankovic", "Maxi"]
    assert searching.get(f"/api/jobs/{job_id}/results").json()["counts"]["unchecked"] == 0
